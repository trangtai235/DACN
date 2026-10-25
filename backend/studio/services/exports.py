"""Interchange formats implemented against a single offset-preserving record schema."""
import csv,io,json,re,xml.etree.ElementTree as ET
from collections import Counter
from openpyxl import Workbook
from studio.models import Annotation,GoldAnnotation
from studio.services.projects import labels_for,ensure
FORMATS={
 'conll':('CoNLL BIO','conll','text/plain'), 'conll_full':('CoNLL offsets','tsv','text/tab-separated-values'),
 'hf_jsonl':('Hugging Face NER','jsonl','application/x-ndjson'), 'spacy_json':('spaCy JSON','json','application/json'),
 'doccano_jsonl':('Doccano','jsonl','application/x-ndjson'), 'label_studio_json':('Label Studio','json','application/json'),
 'json_sentence':('JSON câu','json','application/json'), 'json_token':('JSON token','json','application/json'),
 'jsonl':('JSONL đầy đủ','jsonl','application/x-ndjson'), 'xml':('XML corpus','xml','application/xml'),
 'csv_sentence':('CSV câu','csv','text/csv'), 'csv_token':('CSV token','csv','text/csv'),
 'csv_spans':('CSV spans','csv','text/csv'), 'csv_annotations':('CSV người gán','csv','text/csv'),
 'xlsx':('Excel 3 trang','xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
 'json_llm':('LLM chat fine-tuning','jsonl','application/x-ndjson')}
def protect(value):
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):return "'"+value
    return value

def tokens(text,spans):
    # Split whitespace tokens at every accepted span boundary, preserving partial-token labels.
    result=[];boundaries={s[k] for s in spans for k in ('start','end')}
    for m in re.finditer(r'\S+',text):
        points=sorted({m.start(),m.end()}|{p for p in boundaries if m.start()<p<m.end()})
        for start,end in zip(points,points[1:]):
            span=next((s for s in spans if s['start']<=start and s['end']>=end),None)
            group=spans.index(span) if span else None;key=span['label'] if span else 'O'
            bio=(('I-' if result and result[-1]['span_group']==group else 'B-')+key) if span else 'O'
            result.append({'text':text[start:end],'start':start,'end':end,'position':len(result),'label':key,'bio':bio,'span_group':group})
    return result

def records(project,scope='all',source=None):
    if scope not in ('all','labelled','human','gold'):raise ValueError('Phạm vi xuất không hợp lệ.')
    prefs=ensure(project);qs=project.comments.select_related('source').order_by('id')
    if source:qs=qs.filter(source_id=source)
    out=[]
    for c in qs.prefetch_related('annotations__label','annotations__user','predictions'):
        anns=list(c.annotations.all());submitted=[a for a in anns if a.status=='submitted'];keys=[a.label.key if a.label else 'O' for a in submitted]
        sentence_agreed=len(keys)>=prefs.redundancy and len(set(keys))==1
        span_signatures={json.dumps(a.spans,sort_keys=True) for a in submitted}
        agreed=sentence_agreed and len(span_signatures)==1
        gold=GoldAnnotation.objects.filter(comment=c).select_related('label').first()
        manual=sorted(submitted,key=lambda a:a.updated_at)[-1] if submitted else None
        latest=c.predictions.order_by('-id').first()
        key=(gold.label.key if gold.label else 'O') if gold else (manual.label.key if manual and manual.label else 'O') if manual else c.first_label or c.ai_label or 'O'
        spans=gold.spans if gold else manual.spans if manual else latest.spans if latest else []
        human=bool(gold or manual or c.first_label);status='adjudicated' if gold else 'agreed' if agreed else 'conflict' if len(set(keys))>1 else 'annotated' if human else 'ai' if c.ai_label else 'pending'
        if scope=='labelled' and key=='O' and not human:continue
        if scope=='human' and not human:continue
        if scope=='gold' and not (gold or agreed):continue
        row={'id':c.id,'text':c.text,'source_text':c.original_text or c.text,'context':c.context,'source':c.source.name,'source_id':c.source_id,'external_id':c.external_id,'author':c.author,'parent_external_id':c.parent_external_id,'label':key,'ai_label':c.ai_label,'manual_label':manual.label.key if manual and manual.label else c.first_label,'gold_label':(gold.label.key if gold.label else 'O') if gold else keys[0] if agreed else None,'review_status':status,'is_meaningful':not c.unusable,'annotator_count':len(submitted),'annotators':[{'annotator':a.user.username,'user_id':a.user_id,'source':'human','label':a.label.key if a.label else None,'spans':a.spans,'note':a.note,'assisted':a.assisted,'seconds':a.seconds,'status':a.status,'revision':a.revision,'created_at':a.created_at.isoformat()} for a in anns],'spans':[{**s,'text':c.text[s['start']:s['end']]} for s in spans]}
        row['tokens']=tokens(c.text,spans);out.append(row)
    return out

def render(project,rows,fmt):
    if fmt not in FORMATS:raise ValueError('Định dạng xuất không hợp lệ.')
    title,ext,mime=FORMATS[fmt];j=lambda x:json.dumps(x,ensure_ascii=False)
    if fmt in ('json_sentence','json_token'):
        body=j([{k:v for k,v in r.items() if fmt=='json_token' or k not in ('tokens','spans')} for r in rows]).encode()
    elif fmt=='jsonl':body=('\n'.join(j(r) for r in rows)+'\n').encode()
    elif fmt in ('conll','conll_full'):
        parts=[]
        for r in rows:
            parts.append(f"# id = {r['id']}\n# label = {r['label']}")
            parts.extend('\t'.join(str(v) for v in ([t['text'],t['bio']] if fmt=='conll' else [t['text'],t['start'],t['end'],t['bio'],r['label']])) for t in r['tokens']);parts.append('')
        body='\n'.join(parts).encode()
    elif fmt=='hf_jsonl':
        keys=[x['key'] for x in labels_for(project,False)];names=['O']+sorted([p+k for k in keys for p in ('B-','I-')]);index={k:i for i,k in enumerate(names)}
        body=('\n'.join([j({'__meta__':True,'project':project.name,'label_names':names})]+[j({'id':r['id'],'tokens':[t['text'] for t in r['tokens']],'ner_tags':[index.get(t['bio'],0) for t in r['tokens']],'bio_tags':[t['bio'] for t in r['tokens']],'sentence_label':r['label'],'review_status':r['review_status']}) for r in rows])+'\n').encode()
    elif fmt=='spacy_json':body=j([[r['text'],{'entities':[[s['start'],s['end'],s['label']] for s in r['spans']],'cats':{r['label']:1.0} if r['label']!='O' else {}}] for r in rows]).encode()
    elif fmt=='doccano_jsonl':body=('\n'.join(j({'id':r['id'],'text':r['text'],'label':[[s['start'],s['end'],s['label']] for s in r['spans']],'cats':[r['label']] if r['label']!='O' else []}) for r in rows)+'\n').encode()
    elif fmt=='label_studio_json':
        body=j([{'data':{'text':r['text'],'ref_id':r['id']},'annotations':[{'result':[{'from_name':'label','to_name':'text','type':'labels','value':{'start':s['start'],'end':s['end'],'text':s['text'],'labels':[s['label']]}} for s in r['spans']]+[{'from_name':'sentence','to_name':'text','type':'choices','value':{'choices':[r['label']]}}]}]} for r in rows]).encode()
    elif fmt=='json_llm':body=('\n'.join(j({'messages':[{'role':'system','content':'Gán nhãn bình luận tiếng Việt và spans; trả JSON.'},{'role':'user','content':r['text']},{'role':'assistant','content':j({'label':r['label'],'is_meaningful':r['is_meaningful'],'spans':r['spans']})}],'metadata':{'id':r['id'],'review_status':r['review_status']}}) for r in rows)+'\n').encode()
    elif fmt=='xml':
        root=ET.Element('corpus',project=project.name)
        for r in rows:
            node=ET.SubElement(root,'sentence',id=str(r['id']),label=r['label'],review_status=r['review_status']);ET.SubElement(node,'text').text=r['text'];ts=ET.SubElement(node,'tokens')
            for t in r['tokens']:ET.SubElement(ts,'token',start=str(t['start']),end=str(t['end']),bio=t['bio']).text=t['text']
            ss=ET.SubElement(node,'spans')
            for s in r['spans']:ET.SubElement(ss,'span',start=str(s['start']),end=str(s['end']),label=s['label']).text=s['text']
        body=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    else:
        fields=['id','text','source','label','ai_label','manual_label','gold_label','review_status','annotator_count','is_meaningful','author'];data=[]
        if fmt in ('csv_sentence','xlsx'):data=[{k:r.get(k) for k in fields} for r in rows]
        elif fmt=='csv_token':
            fields=['comment_id','comment_label','token_text','token_label','bio','position','start','end','span_group'];data=[{'comment_id':r['id'],'comment_label':r['label'],'token_text':t['text'],'token_label':t['label'],**{k:t[k] for k in fields if k in t}} for r in rows for t in r['tokens']]
        elif fmt=='csv_spans':
            fields=['comment_id','comment_text','comment_label','span_text','span_label','start','end'];data=[{'comment_id':r['id'],'comment_text':r['text'],'comment_label':r['label'],'span_text':s['text'],'span_label':s['label'],'start':s['start'],'end':s['end']} for r in rows for s in r['spans']]
        elif fmt=='csv_annotations':
            fields=['comment_id','text','annotator','source','label','assisted','created_at'];data=[{'comment_id':r['id'],'text':r['text'],**{k:a.get(k) for k in fields if k in a}} for r in rows for a in r['annotators']]
        if fmt=='xlsx':
            book=Workbook();sheet=book.active;sheet.title='Câu';sheet.append(fields)
            for row in data:sheet.append([protect(row.get(k)) for k in fields])
            spans=book.create_sheet('Cụm đã gán nhãn');spans.append(['id','text','label','start','end'])
            for r in rows:
                for s in r['spans']:spans.append([r['id'],protect(s['text']),s['label'],s['start'],s['end']])
            meta=book.create_sheet('Thông tin');meta.append(['Project',protect(project.name)]);meta.append(['Records',len(rows)]);meta.append(['Offsets','Unicode code points, end exclusive'])
            buf=io.BytesIO();book.save(buf);body=buf.getvalue()
        else:
            buf=io.StringIO();writer=csv.DictWriter(buf,fieldnames=fields);writer.writeheader()
            for row in data:writer.writerow({k:protect(v) for k,v in row.items()})
            body=('\ufeff'+buf.getvalue()).encode()
    return body,ext,mime
