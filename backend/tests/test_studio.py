import hashlib,io,json,tempfile
from pathlib import Path
from unittest.mock import patch
from django.test import TestCase,Client,override_settings
from django.contrib.auth.models import User
from django.core import mail
from django.utils import timezone
from core.models import Project,Membership,Source,Comment,Job,Prediction
from studio.models import Annotation,AnnotationEvent,GoldAnnotation,DatasetVersion,ProjectLabel,AIExposure,UserProfile
from studio.services.projects import ensure,labels_for
from studio.services.quality import quality
from studio.services.exports import records,render,tokens,FORMATS
from integrations.models import ProviderConnection
from integrations.secrets import encrypt,decrypt
from integrations.providers.base import parse_result,validate_endpoint
from integrations.providers.ollama import OllamaAdapter
from core.jobs.runner import execute,tick,Cancelled
from experiments.models import ModelRun
from django.conf import settings

class StudioTests(TestCase):
    def setUp(self):
        self.users=[User.objects.create_user(username=f'studio{i}',password='Test-Pass!47',email=f'studio{i}@example.test') for i in range(4)]
        self.project=Project.objects.create(name='Studio project',owner=self.users[0]);self.pref=ensure(self.project)
        for i,u in enumerate(self.users[:3]):Membership.objects.create(project=self.project,user=u,role='manager' if i==0 else 'annotator')
        self.source=Source.objects.create(project=self.project,name='Test source')
        self.comment=Comment.objects.create(project=self.project,source=self.source,text='🙂 bạn ngu',ai_label='TOXIC')
        self.client.force_login(self.users[0]);self.label=ProjectLabel.objects.get(project=self.project,key='TOXIC')
    def post(self,path,p):return self.client.post('/api/'+path,json.dumps(p),content_type='application/json')
    def payload(self,**kwargs):return {'project':self.project.id,**kwargs}
    def test_delete_annotated_source(self):
        from reviewing.models import ReviewRun,ReviewTask
        from core.models import AuditEvent
        self.submit()
        GoldAnnotation.objects.create(comment=self.comment,label=self.label,adjudicator=self.users[0])
        self.comment.first_frozen_at=timezone.now();self.comment.save()
        run=ReviewRun.objects.create(project=self.project)
        task=ReviewTask.objects.create(run=run,comment=self.comment,reviewer=self.users[1])
        other=Source.objects.create(project=self.project,name='Keep')
        kept=Comment.objects.create(project=self.project,source=other,text='Keep this comment')
        path=f'studio/sources/{self.source.id}/action'
        self.assertEqual(self.post(path,{'action':'delete','confirm':'wrong'}).status_code,400)
        job=Job.objects.create(project=self.project,actor=self.users[0],kind='predict')
        self.assertEqual(self.post(path,{'action':'delete','confirm':self.source.name}).status_code,400)
        self.assertTrue(ReviewTask.objects.filter(pk=task.pk).exists())
        job.status='completed';job.save()
        self.client.force_login(self.users[1])
        self.assertNotEqual(self.post(path,{'action':'delete','confirm':self.source.name}).status_code,200)
        self.client.force_login(self.users[0])
        self.assertEqual(self.post(path,{'action':'delete','confirm':self.source.name}).status_code,200)
        self.assertFalse(Source.objects.filter(pk=self.source.pk).exists())
        self.assertFalse(Comment.objects.filter(pk=self.comment.pk).exists())
        self.assertFalse(Annotation.objects.filter(comment_id=self.comment.pk).exists())
        self.assertFalse(GoldAnnotation.objects.filter(comment_id=self.comment.pk).exists())
        self.assertFalse(ReviewTask.objects.filter(pk=task.pk).exists())
        self.assertTrue(Comment.objects.filter(pk=kept.pk).exists())
        self.assertTrue(AuditEvent.objects.filter(project=self.project,action='source.delete').exists())
    def submit(self,user=0,**kw):
        self.client.force_login(self.users[user]);return self.post(f'studio/comments/{self.comment.id}/annotation',self.payload(label='TOXIC',spans=[{'start':6,'end':9,'label':'TOXIC'}],revision=0,seconds=10,**kw))
    def test_workspace_source_queue_stays_in_selected_source(self):
        other=Source.objects.create(project=self.project,name='Other video')
        second=Comment.objects.create(project=self.project,source=other,text='Another comment')
        response=self.client.get('/api/studio/workspace',{'project':self.project.id,'source':other.id})
        self.assertEqual(response.json()['comment']['id'],second.id)
        saved=self.post(f'studio/comments/{second.id}/annotation',self.payload(label='TOXIC',spans=[],revision=0))
        self.assertEqual(saved.status_code,200)
        response=self.client.get('/api/studio/workspace',{'project':self.project.id,'source':other.id})
        self.assertIsNone(response.json()['comment'])
        response=self.client.get('/api/studio/workspace',{'project':self.project.id})
        self.assertEqual(response.json()['comment']['id'],self.comment.id)

    def test_workspace_exposes_youtube_comment_origin(self):
        self.source.kind='youtube';self.source.external_id='video123'
        self.source.save()
        self.comment.author='@viewer';self.comment.external_id='comment123'
        self.comment.save()
        response=self.client.get('/api/studio/workspace',{'project':self.project.id,'comment':self.comment.id})
        self.assertEqual(response.status_code,200)
        row=response.json()['comment']
        self.assertEqual(row['source_kind'],'youtube')
        self.assertEqual(row['video_id'],'video123')
        self.assertEqual(row['external_id'],'comment123')
        self.assertEqual(row['author'],'@viewer')
        self.assertIsNone(row['published_at'])

    def test_unicode_spans_and_stale_revision(self):
        r=self.submit();self.assertEqual(r.status_code,200);self.assertEqual(r.json()['annotation']['spans'][0]['start'],6)
        r=self.submit();self.assertEqual(r.status_code,400);self.assertEqual(Annotation.objects.count(),1)
        self.assertEqual(self.comment.text[6:9],'ngu')
    def test_invalid_overlap_and_foreign_label_rejected(self):
        r=self.post(f'studio/comments/{self.comment.id}/annotation',self.payload(label='FOREIGN',spans=[],revision=0));self.assertEqual(r.status_code,400)
        r=self.post(f'studio/comments/{self.comment.id}/annotation',self.payload(label='TOXIC',spans=[{'start':1,'end':5,'label':'TOXIC'},{'start':3,'end':6,'label':'TOXIC'}],revision=0));self.assertEqual(r.status_code,400)
        self.assertFalse(Annotation.objects.exists())
    def test_reviewed_no_label_is_distinct_from_missing(self):
        r=self.post(f'studio/comments/{self.comment.id}/annotation',self.payload(label=None,spans=[],revision=0));self.assertEqual(r.status_code,200)
        self.assertEqual(Annotation.objects.get().status,'submitted');self.assertIsNone(Annotation.objects.get().label)
        result=self.client.get('/api/studio/comments',{'project':self.project.id,'status':'done'}).json();self.assertEqual(result['total'],1)
    def test_redundancy_and_skip_queue(self):
        self.assertEqual(self.submit(0).status_code,200);self.assertEqual(self.submit(1).status_code,200);self.assertEqual(self.submit(2).status_code,400)
        self.client.force_login(self.users[2]);result=self.client.get('/api/studio/workspace',{'project':self.project.id}).json();self.assertIsNone(result['comment'])
    def test_private_workspace_and_non_member(self):
        result=self.client.get('/api/studio/workspace',{'project':self.project.id}).json();self.assertTrue(result['comment']['ai_hidden']);self.assertNotIn('ai_label',result['comment'])
        self.client.force_login(self.users[3]);self.assertEqual(self.client.get('/api/studio/workspace',{'project':self.project.id}).status_code,404)
    def test_exposure_is_sticky_on_revision(self):
        self.submit();self.client.get('/api/studio/workspace',{'project':self.project.id,'comment':self.comment.id});self.assertTrue(AIExposure.objects.exists())
        r=self.post(f'studio/comments/{self.comment.id}/annotation',self.payload(label='NON_TOXIC',spans=[],revision=1));self.assertEqual(r.status_code,200);self.assertTrue(Annotation.objects.get().assisted)
    def test_lock_prevents_import_annotation_and_ai(self):
        self.pref.locked=True;self.pref.save()
        for path,body in [('import',self.payload(text='test')),('predict',self.payload(provider='local')),(f'studio/comments/{self.comment.id}/annotation',self.payload(label='TOXIC',revision=0,spans=[]))]:self.assertEqual(self.post(path,body).status_code,403)
        self.assertEqual(self.post('studio/project-settings',self.payload(locked=False)).status_code,200)
        self.assertEqual(self.submit().status_code,200)
    def test_labels_used_only_in_spans_cannot_be_deleted(self):
        extra=ProjectLabel.objects.create(project=self.project,key='SPAN_ONLY',name='Span only')
        Annotation.objects.create(comment=self.comment,user=self.users[0],label=self.label,spans=[{'start':6,'end':9,'label':'SPAN_ONLY'}])
        r=self.post('studio/labels',self.payload(key='SPAN_ONLY',name='Span only',action='delete',color='#123456'));self.assertEqual(r.status_code,400);self.assertTrue(ProjectLabel.objects.filter(pk=extra.id).exists())
    def test_undo_is_one_reversible_edit(self):
        self.submit();self.assertEqual(self.post(f'studio/comments/{self.comment.id}/undo',self.payload()).status_code,200);self.assertFalse(Annotation.objects.exists())
        self.assertEqual(self.post(f'studio/comments/{self.comment.id}/undo',self.payload()).status_code,400);self.comment.refresh_from_db();self.assertIsNone(self.comment.first_label)
    def test_quality_null_and_disagreement(self):
        self.submit();self.client.force_login(self.users[1]);self.post(f'studio/comments/{self.comment.id}/annotation',self.payload(label=None,spans=[],revision=0))
        r=quality(self.project);self.assertEqual(r['overlap_units'],1);self.assertEqual(r['raw_agreement'],0);self.assertIn(self.comment.id,r['conflict_ids']);self.assertIsNotNone(r['krippendorff_alpha'])
    def test_span_disagreement_not_gold(self):
        self.submit();Annotation.objects.create(comment=self.comment,user=self.users[1],label=self.label,spans=[])
        self.assertIn(self.comment.id,quality(self.project)['conflict_ids']);self.assertEqual(records(self.project,'gold'),[])
    def test_gold_no_label_and_human_scope(self):
        r=self.post(f'studio/comments/{self.comment.id}/gold',self.payload(label=None,spans=[],note='No sentence label appropriate'));self.assertEqual(r.status_code,200)
        self.assertEqual(records(self.project,'human')[0]['label'],'O')
    def test_versions_restore_monotonic_revision_backup_and_null_gold(self):
        self.submit();GoldAnnotation.objects.create(comment=self.comment,label=None,note='No label',adjudicator=self.users[0])
        r=self.post('studio/versions',self.payload(name='v1'));self.assertEqual(r.status_code,200);v=DatasetVersion.objects.get(pk=r.json()['id'])
        self.assertEqual(v.sha256,hashlib.sha256(json.dumps(v.snapshot,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest())
        a=Annotation.objects.get();a.revision=9;a.seconds=999;a.save()
        r=self.post(f'studio/versions/{v.id}/restore',self.payload(confirm='v1'));self.assertEqual(r.status_code,200);a=Annotation.objects.get();self.assertGreater(a.revision,9);self.assertEqual(a.seconds,10);self.assertIsNone(GoldAnnotation.objects.get().label);self.assertEqual(DatasetVersion.objects.count(),3)
    def test_export_all_sixteen_formats_and_offsets(self):
        self.submit();rows=records(self.project,'human');self.assertEqual(rows[0]['source_text'],self.comment.text)
        self.assertEqual(len(FORMATS),16)
        from openpyxl import load_workbook
        for fmt in FORMATS:
            content,ext,mime=render(self.project,rows,fmt);self.assertTrue(content,fmt)
            if fmt=='xlsx':self.assertEqual(len(load_workbook(io.BytesIO(content)).sheetnames),3)
        self.assertEqual([t['bio'] for t in rows[0]['tokens']],['O','O','B-TOXIC'])
    def test_partial_token_spans_not_silently_lost(self):
        t=tokens('(ngu)',[{'start':1,'end':4,'label':'TOXIC'}]);self.assertEqual([x['text'] for x in t],['(','ngu',')']);self.assertEqual(t[1]['bio'],'B-TOXIC')
    def test_csv_and_xlsx_formula_protection(self):
        self.comment.text='=1+1';self.comment.save();rows=records(self.project,'all');data,_,_=render(self.project,rows,'csv_sentence');self.assertIn("'=1+1",data.decode())
        from openpyxl import load_workbook
        data,_,_=render(self.project,rows,'xlsx');value=load_workbook(io.BytesIO(data)).worksheets[0]['B2'];self.assertNotEqual(value.data_type,'f')
    def test_provider_secret_preserved_and_not_returned(self):
        body=self.payload(name='Local Ollama',kind='ollama',base_url='http://127.0.0.1:11434',model='test-model',api_key='secret-not-in-response')
        r=self.post('studio/connections/save',body);self.assertEqual(r.status_code,200);self.assertNotIn('secret-not-in-response',r.content.decode());c=ProviderConnection.objects.get();self.assertNotEqual(c.encrypted_key,body['api_key']);self.assertEqual(decrypt(c.encrypted_key),body['api_key'])
        body.update(id=c.id,api_key='');self.post('studio/connections/save',body);c.refresh_from_db();self.assertEqual(decrypt(c.encrypted_key),'secret-not-in-response')
        result=self.client.get('/api/studio/connections',{'project':self.project.id});self.assertNotIn('secret-not-in-response',result.content.decode())
    def test_endpoint_allowlist_and_provider_validation(self):
        with self.assertRaises(ValueError):validate_endpoint('http://169.254.169.254/latest/meta-data')
        with self.assertRaises(ValueError):validate_endpoint('https://api.openai.com.evil.test/v1')
        with self.assertRaises(ValueError):parse_result('{"label":"BAD"}',self.comment.text,labels_for(self.project))
        with self.assertRaises(ValueError):parse_result({'label':'TOXIC','spans':['bad']},self.comment.text,labels_for(self.project))
    def test_ollama_adapter_structured_response(self):
        connection=ProviderConnection(project=self.project,name='Test',kind='ollama',base_url='http://127.0.0.1:11434',model='test-model')
        with patch('integrations.providers.ollama.httpx.Client') as client:
            response=client.return_value.__enter__.return_value.post.return_value;response.status_code=200;response.json.return_value={'message':{'content':'{"label":"TOXIC","spans":[{"start":6,"end":9,"label":"TOXIC"}]}'},'prompt_eval_count':10,'eval_count':5}
            result=OllamaAdapter().predict(self.comment.text,labels_for(self.project),connection,'');self.assertEqual(result['label'],'TOXIC');self.assertEqual(result['input_tokens'],10);self.assertEqual(result['spans'][0]['start'],6)
    def test_invite_matches_email_and_consumes_once(self):
        r=self.post('studio/invite',self.payload(email=self.users[3].email,send_email=False));self.assertEqual(r.status_code,200);token=r.json()['url'].split('invite=')[1]
        self.client.force_login(self.users[1]);self.assertEqual(self.post('studio/accept-invite',{'token':token}).status_code,403)
        self.client.force_login(self.users[3]);self.assertEqual(self.post('studio/accept-invite',{'token':token}).status_code,200);self.assertEqual(self.post('studio/accept-invite',{'token':token}).status_code,400)
    @override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
    def test_verification_email_and_password_reset(self):
        r=self.post('studio/verify/send',{});self.assertEqual(r.status_code,200);self.assertEqual(len(mail.outbox),1);token=mail.outbox[0].body.split('#verify=')[1]
        self.assertEqual(self.post('studio/verify',{'token':token}).status_code,200);self.assertTrue(UserProfile.objects.get(user=self.users[0]).email_verified)
        self.client.logout();r=self.post('studio/password/request',{'email':self.users[0].email});self.assertEqual(r.status_code,200);uid,token=mail.outbox[-1].body.split('#reset=')[1].split('~');body={'user_id':uid,'token':token,'password':'Fresh-Password!59','password_confirm':'Fresh-Password!59'}
        self.assertEqual(self.post('studio/password/reset',body).status_code,200);self.assertEqual(self.post('studio/password/reset',body).status_code,400);self.users[0].refresh_from_db();self.assertTrue(self.users[0].check_password('Fresh-Password!59'))
    def test_json_list_rejected_as_client_error(self):
        response=self.client.post('/api/studio/selection','[]',content_type='application/json');self.assertEqual(response.status_code,400)
    def test_real_local_prediction_job_without_redis(self):
        root=settings.BASE_DIR/'artifacts/victsd-logreg-v1';ModelRun.objects.create(project=self.project,name='Baseline',artifact_path=str(root),report=(root/'metrics.json').read_text())
        job=Job.objects.create(project=self.project,actor=self.users[0],kind='predict',payload=json.dumps({'provider':'local','replace_ai':True}))
        execute(job.id);job.refresh_from_db();self.assertEqual(job.status,'complete');self.assertEqual(Prediction.objects.count(),1);self.assertIn('identity',job.checkpoint)
    def test_local_prediction_saves_batches_before_more_inference(self):
        from core.jobs.predict import run
        from types import SimpleNamespace
        for i in range(9):
            Comment.objects.create(project=self.project,source=self.source,text=f'batch {i}',digest=hashlib.sha256(f'batch {i}'.encode()).hexdigest())
        model=SimpleNamespace(id=123,algorithm='transformer_bio',artifact_path='unused')
        job=Job.objects.create(project=self.project,actor=self.users[0],kind='predict')
        sizes=[]
        def infer(project,texts,model_id):
            if sizes:self.assertEqual(Prediction.objects.count(),8)
            sizes.append(len(texts))
            return [{'label':'TOXIC','spans':[],'model_version':'test'} for text in texts]
        with patch('core.jobs.predict.model_for',return_value=model), patch('core.jobs.predict.artifact_digest',return_value='test'), patch('core.jobs.predict.local_predict',side_effect=infer):
            run(job,{'provider':'local','replace_ai':True})
        self.assertEqual(sizes,[8,2])
        self.assertEqual(Prediction.objects.count(),10)

    def test_completed_checkpoint_survives_cancellation_and_revocation(self):
        job=Job.objects.create(project=self.project,actor=self.users[0],kind='predict',cancel_requested=True)
        with self.assertRaises(Cancelled):tick(job,50,'Completed response',{'predictions':[{'label':'TOXIC'}]})
        job.refresh_from_db();self.assertEqual(len(job.checkpoint['predictions']),1)
        job.cancel_requested=False;job.save();Membership.objects.filter(project=self.project,user=self.users[0]).delete()
        with self.assertRaises(Cancelled):tick(job,50,'No access')
    def test_benchmark_configuration_drift_is_rejected(self):
        from core.jobs.pinning import pin
        job=Job.objects.create(project=self.project,actor=self.users[0],kind='benchmark_api');pin(job,{'model':'one'})
        with self.assertRaises(ValueError):pin(job,{'model':'two'})
    @override_settings(YOUTUBE_API_KEY='test-only-key')
    def test_youtube_cursor_resumes_inside_reply_page(self):
        from core.jobs.youtube import run
        source=Source.objects.create(project=self.project,name='YouTube',kind='youtube',metadata={'title':'Fixture'})
        def item(i):return {'id':i,'snippet':{'textOriginal':i+' text','authorDisplayName':'Fixture'}}
        def response(url,params):
            from types import SimpleNamespace
            if url.endswith('commentThreads'):data={'items':[{'snippet':{'topLevelComment':item('root'),'totalReplyCount':3}},{'snippet':{'topLevelComment':item('root2'),'totalReplyCount':0}}]}
            elif params.get('pageToken')=='reply2':data={'items':[item('reply3')]}
            else:data={'items':[item('reply1'),item('reply2')],'nextPageToken':'reply2'}
            return SimpleNamespace(status_code=200,json=lambda:data)
        job=Job.objects.create(project=self.project,actor=self.users[0],kind='youtube',payload='{}')
        payload={'source_id':source.id,'video_id':'testvideo01','limit':2,'include_replies':True}
        with patch('core.jobs.youtube.httpx.Client') as client:
            client.return_value.__enter__.return_value.get.side_effect=response
            first=run(job,payload);self.assertEqual(first['imported_total'],2);source.refresh_from_db();self.assertEqual(source.fetch_state,'paused');self.assertTrue(source.checkpoint['root_done']);self.assertEqual(source.checkpoint['reply_index'],1)
            retry=Job.objects.create(project=self.project,actor=self.users[0],kind='youtube',checkpoint=source.checkpoint)
            second=run(retry,{**payload,'limit':None});self.assertTrue(second['complete']);self.assertEqual(second['imported_total'],5)
        self.assertEqual(source.comments.count(),5);self.assertEqual(source.comments.get(external_id='reply3').context,'root text')
    @override_settings(YOUTUBE_API_KEY='test-only-key')
    def test_youtube_lock_during_request_blocks_write(self):
        from core.jobs.youtube import run
        source=Source.objects.create(project=self.project,name='YT',kind='youtube',metadata={'title':'Fixture'})
        job=Job.objects.create(project=self.project,actor=self.users[0],kind='youtube')
        def response(*args,**kwargs):
            from types import SimpleNamespace
            self.pref.locked=True;self.pref.save()
            return SimpleNamespace(status_code=200,json=lambda:{'items':[{'snippet':{'topLevelComment':{'id':'root','snippet':{'textOriginal':'must not persist'}},'totalReplyCount':0}}]})
        with patch('core.jobs.youtube.httpx.Client') as client:
            client.return_value.__enter__.return_value.get.side_effect=response
            with self.assertRaises(PermissionError):run(job,{'source_id':source.id,'video_id':'testvideo01','limit':10})
        self.assertEqual(source.comments.count(),0)
    @override_settings(YOUTUBE_API_KEY='test-only-key')
    def test_youtube_unlimited_collects_all_pages(self):
        from core.jobs.youtube import run
        from types import SimpleNamespace
        source=Source.objects.create(project=self.project,name='YT',kind='youtube',metadata={'title':'Fixture'})
        job=Job.objects.create(project=self.project,actor=self.users[0],kind='youtube',checkpoint={'count':10000})
        def response(url,params):
            token=params.get('pageToken','first')
            data={'items':[{'snippet':{'topLevelComment':{'id':token,'snippet':{'textOriginal':token}},'totalReplyCount':0}}]}
            if token=='first':data['nextPageToken']='second'
            return SimpleNamespace(status_code=200,json=lambda:data)
        with patch('core.jobs.youtube.httpx.Client') as client:
            client.return_value.__enter__.return_value.get.side_effect=response
            result=run(job,{'source_id':source.id,'video_id':'testvideo01','limit':''})
        self.assertTrue(result['complete'])
        self.assertEqual(result['imported_total'],10002)
        self.assertEqual(source.comments.count(),2)

    def test_collection_limit_validation(self):
        from core.services.collection import collection_limit
        for value in (None,'',' '):self.assertIsNone(collection_limit(value))
        self.assertEqual(collection_limit('15000'),15000)
        for value in (0,-1,'abc','1.5',True):
            with self.assertRaises(ValueError):collection_limit(value)

    def test_model_and_api_selection_are_project_scoped(self):
        other=Project.objects.create(name='Other',owner=self.users[3]);ensure(other)
        model=ModelRun.objects.create(project=other,name='Foreign',artifact_path='none',report='{}')
        self.assertEqual(self.post('studio/selection',self.payload(model_id=model.id)).status_code,404)
        self.client.force_login(self.users[1]);self.assertEqual(self.post('studio/selection',self.payload(model_id=model.id)).status_code,403)
    def test_training_three_classes_and_all_algorithms(self):
        import csv
        from experiments.training import train
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);paths=[];keys=['CLEAN','OFFENSIVE','HATE']
            for split in ('train','validation','test'):
                path=root/(split+'.csv');paths.append(path)
                with path.open('w',newline='',encoding='utf-8') as f:
                    writer=csv.writer(f);writer.writerow(['text','label'])
                    for i,key in enumerate(keys):
                        for j in range(4):writer.writerow([f'{split} class{i} repeated feature{i} example {j}',key])
            for algorithm in ('logistic_regression','linear_svm','complement_nb'):
                report=train(*paths,root/algorithm,algorithm_key=algorithm,label_keys=keys)
                self.assertEqual(report['label_keys'],keys);self.assertEqual(len(report['test']['confusion_matrix']),3);self.assertEqual(report['test']['n'],12)
    def test_native_seed_is_idempotent_and_registers_three_models(self):
        from django.core.management import call_command
        call_command('seed_demo',password='Synthetic-Pass!47',stdout=io.StringIO())
        call_command('seed_demo',password='Unused-Pass!48',stdout=io.StringIO())
        demo=Project.objects.get(is_demo=True);self.assertEqual(demo.model_runs.count(),3) if hasattr(demo,'model_runs') else self.assertEqual(ModelRun.objects.filter(project=demo).count(),3)
        self.assertTrue(Annotation.objects.filter(comment__project=demo).exists());self.assertTrue(User.objects.get(username='tv1').check_password('Synthetic-Pass!47'))
    def test_skipping_preserves_hidden_ai(self):
        self.assertEqual(self.post(f'studio/comments/{self.comment.id}/annotation',self.payload(status='skipped',revision=0)).status_code,200)
        result=self.client.get('/api/studio/workspace',{'project':self.project.id,'comment':self.comment.id}).json()['comment']
        self.assertTrue(result['ai_hidden']);self.assertNotIn('ai_label',result)
    def test_invitation_list_and_revoke(self):
        r=self.post('studio/invite',self.payload(email=self.users[3].email,send_email=False));self.assertEqual(r.status_code,200);value=r.json()
        listing=self.client.get('/api/studio/invite',{'project':self.project.id});self.assertEqual(listing.status_code,200);self.assertEqual(listing.json()['invitations'][0]['id'],value['id'])
        self.assertEqual(self.post('studio/invite',self.payload(id=value['id'],action='revoke')).status_code,200)
        self.client.force_login(self.users[3]);self.assertEqual(self.post('studio/accept-invite',{'token':value['url'].split('invite=')[1]}).status_code,400)
    def test_blind_task_cannot_reveal_first_pass_api(self):
        from reviewing.models import ReviewRun,ReviewTask
        run=ReviewRun.objects.create(project=self.project,policy='random',population='[]')
        ReviewTask.objects.create(run=run,comment=self.comment,reviewer=self.users[0])
        r=self.post(f'comments/{self.comment.id}/first',self.payload(label='TOXIC',revision=self.comment.revision));self.assertEqual(r.status_code,403)
